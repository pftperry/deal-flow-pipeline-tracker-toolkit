"""Builds the Deal Flow Pipeline Tracker Toolkit workbook (Post Fiat Analytical Infrastructure Library).
Usage: python build_deal_flow_tracker.py [out.xlsx]"""
import sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1] if len(sys.argv) > 1 else "Deal_Flow_Pipeline_Tracker_Toolkit.xlsx"
N = ["ReadMe - Validator Instructions", "Settings and Scoring Rules", "Intermediary Directory",
     "Deal Pipeline Log", "Follow-Up Cadence Tracker", "Ranking Dashboard", "Anonymized Worked Example"]
wb = Workbook(); wb.active.title = N[0]
for n in N[1:]:
    wb.create_sheet(n)
RM, ST, DR, PL, FT, DB, WE = (wb[n] for n in N)
S, D, P, T, B, W = (f"'{n}'!" for n in N[1:])
DR0, DR1, PL0, PL1 = 4, 28, 4, 43          # directory rows 4-28, pipeline rows 4-43

thin = Side(style="thin", color="BFBFBF"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {"in": PatternFill("solid", fgColor="FFF2CC"), "f": PatternFill("solid", fgColor="DDEBF7"),
        "h": PatternFill("solid", fgColor="2F5597"), "t": PatternFill("solid", fgColor="1F3864"),
        "k": PatternFill("solid", fgColor="E2EFDA")}
FMT = {"$M": '#,##0.00', "%": '0%', "d": 'yyyy-mm-dd', "n": '0', "n1": '0.0', "x": '0.000'}


def title(ws, text, sub, span):
    for c in range(1, span + 1):
        ws.cell(1, c).fill = FILL["t"]
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")


def header(ws, r, cols):
    for i, v in enumerate(cols, 1):
        c = ws.cell(r, i, v); c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL["h"]
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BOX


def put(ws, r, col, v, fmt=None, kind=None, bold=False):
    c = ws.cell(r, col, v); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind is None and isinstance(v, str) and v.startswith("="):
        kind = "f"
    if kind:
        c.fill = FILL[kind]
    if fmt:
        c.number_format = FMT.get(fmt, fmt)
    if bold:
        c.font = Font(bold=True)
    return c


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def section(ws, r, text):
    ws.cell(r, 1, text).font = Font(bold=True, size=12, color="1F3864")


def dv(ws, src, rng):
    v = DataValidation(type="list", formula1=src, allow_blank=True); ws.add_data_validation(v); v.add(rng)


# ---------------- Settings and Scoring Rules ----------------
title(ST, "Settings and Scoring Rules", "Yellow = input. Every weight, stage probability, cadence and threshold used by the model lives on this tab.", 4)
header(ST, 3, ["Setting", "Value", "Unit", "Note"])
for i, (lab, val, fm, unit, note) in enumerate([
        ("As-of date", date(2026, 10, 4), "d", "date", "Fixed so the worked example is reproducible. For live use, replace with =TODAY()."),
        ("Target EBITDA band - minimum", 3.0, "$M", "$M", "Fund's lower-middle-market sweet spot."),
        ("Target EBITDA band - maximum", 15.0, "$M", "$M", "Deals within 50% outside the band get half credit."),
        ("Due-soon window", 7, "n", "days", "Follow-ups due within this many days are flagged DUE IN Nd."),
        ("Stale deal threshold", 21, "n", "days", "Active deal with no near-term action and no activity for this long is flagged STALE.")]):
    r = 4 + i
    put(ST, r, 1, lab, bold=True); put(ST, r, 2, val, fm, "in"); put(ST, r, 3, unit); put(ST, r, 4, note)
header(ST, 10, ["Priority score component", "Weight", "", "Definition (component scaled 0-1, score = 100 x weighted sum)"])
for i, (lab, w, note) in enumerate([
        ("Stage probability", 0.30, "Probability of close for the deal's stage (stage table below)."),
        ("Thesis fit", 0.30, "Analyst fit score 1-5, scaled (fit - 1) / 4."),
        ("EBITDA size fit", 0.15, "1.0 inside the band; 0.5 within 50% outside the band; 0 otherwise."),
        ("Process quality", 0.15, "Process table below: proprietary deals score highest."),
        ("Intermediary relationship", 0.10, "Relationship strength 1-5 from the directory, scaled (strength - 1) / 4.")]):
    r = 11 + i
    put(ST, r, 1, lab, bold=True); put(ST, r, 2, w, "%", "in"); put(ST, r, 4, note)
put(ST, 16, 1, "Total", bold=True); put(ST, 16, 2, "=SUM(B11:B15)", "%", bold=True)
put(ST, 16, 3, '=IF(ROUND(B16,6)=1,"OK - weights sum to 100%","ERROR - weights must sum to 100%")', kind="k")
header(ST, 18, ["Stage", "Order", "Probability of close", "Active (1/0)"])
stages = [("Sourced", .05, 1), ("NDA / Teaser", .10, 1), ("CIM Review", .20, 1), ("IOI Submitted", .35, 1),
          ("Management Meeting", .50, 1), ("LOI Submitted", .70, 1), ("Exclusivity / Diligence", .85, 1),
          ("Closed", 1.0, 0), ("Passed", 0.0, 0), ("Lost", 0.0, 0)]
for i, (s, p, a) in enumerate(stages):
    r = 19 + i
    put(ST, r, 1, s, kind="in", bold=True); put(ST, r, 2, i + 1); put(ST, r, 3, p, "%", "in"); put(ST, r, 4, a, "n", "in")
header(ST, 30, ["Process type", "Process score (0-1)", "", "Why"])
for i, (pt, sc, why) in enumerate([("Proprietary", 1.0, "No competing bidders; relationship-driven."),
                                   ("Limited auction", 0.6, "Handful of invited buyers."),
                                   ("Broad auction", 0.2, "Price-maximizing process; lowest win rate.")]):
    put(ST, 31 + i, 1, pt, kind="in", bold=True); put(ST, 31 + i, 2, sc, "n1", "in"); put(ST, 31 + i, 4, why)
header(ST, 35, ["Intermediary tier", "Touch cadence (days)", "", "Meaning"])
for i, (tr, cd, why) in enumerate([("A", 14, "Active deal source; every two weeks."), ("B", 30, "Periodic source; monthly."),
                                   ("C", 60, "Occasional source; every two months.")]):
    put(ST, 36 + i, 1, tr, kind="in", bold=True); put(ST, 36 + i, 2, cd, "n", "in"); put(ST, 36 + i, 4, why)
widths(ST, [34, 20, 34, 80])
ASOF, EMIN, EMAX, WIN, STALE = (f"{S}$B${r}" for r in (4, 5, 6, 7, 8))
WST, WFIT, WSZ, WPR, WREL = (f"{S}$B${r}" for r in (11, 12, 13, 14, 15))
STAGES, PROCS, TIERS = f"{S}$A$19:$D$28", f"{S}$A$31:$B$33", f"{S}$A$36:$B$38"

# ---------------- Intermediary Directory ----------------
title(DR, "Intermediary Directory", "Anonymized bankers, brokers and other deal sources. Yellow = input (A-J); K-M are live roll-ups from the Deal Pipeline Log. Rows 4-28.", 13)
header(DR, 3, ["ID", "Code name", "Type", "Firm code", "Region", "Sector focus", "Tier (A/B/C)",
               "Relationship strength (1-5)", "Last contact date", "Notes (anonymized)", "Deals sourced", "Active deals", "Best active priority score"])
inter = [
    ("I01", "Banker A", "Investment bank (LMM)", "Firm 1", "Southeast", "Industrial services", "A", 5, date(2026, 9, 28), "Runs limited auctions; sends early looks."),
    ("I02", "Banker B", "Investment bank (LMM)", "Firm 2", "Midwest", "Business services", "A", 4, date(2026, 9, 10), "Broad auctions; strong deal volume."),
    ("I03", "Broker C", "Business broker", "Firm 3", "Texas", "Specialty distribution", "B", 3, date(2026, 8, 20), "Owner-direct listings; light financials."),
    ("I04", "Broker D", "Business broker", "Firm 4", "Southeast", "Home services", "B", 4, date(2026, 9, 22), "Regional roll-up candidates."),
    ("I05", "Advisor E", "M&A advisory boutique", "Firm 5", "Northeast", "Healthcare services", "A", 3, date(2026, 9, 23), "Larger deals; often above our band."),
    ("I06", "CPA F", "Accounting firm", "Firm 6", "Midwest", "Owner-operator businesses", "C", 4, date(2026, 7, 15), "Succession-planning clients."),
    ("I07", "Attorney G", "Law firm (M&A)", "Firm 7", "Texas", "Industrial services", "C", 3, date(2026, 9, 1), "Estate and succession referrals."),
    ("I08", "Operator H", "Former executive", "n/a", "Southeast", "Industrial services", "B", 5, date(2026, 9, 30), "Operating partner candidate; proprietary intros."),
    ("I09", "Banker I", "Investment bank (LMM)", "Firm 8", "West", "Software-enabled services", "B", 2, date(2026, 6, 30), "Relationship cooled after a passed deal."),
    ("I10", "Founder J", "Founder network", "n/a", "Midwest", "Business services", "C", 4, date(2026, 9, 15), "Peer founders considering exits."),
]
for r in range(DR0, DR1 + 1):
    row = inter[r - DR0] if r - DR0 < len(inter) else [None] * 10
    for c, v in enumerate(row, 1):
        put(DR, r, c, v, "d" if c == 9 else ("n" if c == 8 else None), "in")
    put(DR, r, 11, f'=IF($A{r}="","",COUNTIF({P}$D${PL0}:$D${PL1},$A{r}))', "n")
    put(DR, r, 12, f'=IF($A{r}="","",COUNTIFS({P}$D${PL0}:$D${PL1},$A{r},{P}$N${PL0}:$N${PL1},1))', "n")
    put(DR, r, 13, f'=IF($A{r}="","",IF(L{r}=0,0,_xlfn.MAXIFS({P}$R${PL0}:$R${PL1},{P}$D${PL0}:$D${PL1},$A{r},{P}$N${PL0}:$N${PL1},1)))', "n1")
dv(DR, f"={S}$A$36:$A$38", f"G{DR0}:G{DR1}")
widths(DR, [7, 14, 24, 10, 12, 26, 9, 13, 13, 44, 10, 10, 13]); DR.freeze_panes = "C4"

# ---------------- Deal Pipeline Log ----------------
title(PL, "Deal Pipeline Log", "One row per opportunity. Yellow = input (A-L); blue = live formulas (M-X). Priority score, rank and follow-up flag recalculate from Settings. Rows 4-43.", 24)
header(PL, 3, ["Deal ID", "Code name", "Sector", "Intermediary ID", "Intermediary (lookup)", "Process type", "Stage",
               "LTM EBITDA ($M)", "Thesis fit (1-5)", "Last activity date", "Next action", "Next action due",
               "Stage probability", "Active (1/0)", "EBITDA size fit", "Process score", "Relationship strength",
               "Priority score (0-100)", "Rank key (tie-break)", "Priority rank (active)", "Days since last activity",
               "Days overdue", "Follow-up flag", "Prob.-weighted EBITDA ($M)"])
deals = [
    ("D01", "Project Atlas", "Industrial services", "I01", "Limited auction", "Management Meeting", 6.5, 5, date(2026, 9, 30), "Send diligence request list", date(2026, 10, 8)),
    ("D02", "Project Birch", "Business services", "I02", "Broad auction", "IOI Submitted", 9.0, 3, date(2026, 9, 18), "Follow up on IOI feedback", date(2026, 9, 29)),
    ("D03", "Project Cedar", "Specialty distribution", "I03", "Proprietary", "CIM Review", 4.2, 4, date(2026, 9, 12), "Call owner re: Q3 financials", date(2026, 9, 26)),
    ("D04", "Project Delta", "Home services", "I04", "Limited auction", "NDA / Teaser", 3.1, 4, date(2026, 9, 25), "Sign NDA and request CIM", date(2026, 10, 3)),
    ("D05", "Project Ember", "Healthcare services", "I05", "Broad auction", "CIM Review", 18.0, 3, date(2026, 9, 24), "Decide IOI go / no-go", date(2026, 10, 10)),
    ("D06", "Project Falcon", "Industrial services", "I08", "Proprietary", "LOI Submitted", 7.8, 5, date(2026, 10, 1), "Negotiate LOI exclusivity terms", date(2026, 10, 6)),
    ("D07", "Project Granite", "Owner-operator businesses", "I06", "Proprietary", "Sourced", 2.4, 4, date(2026, 8, 30), "Schedule intro call with owner", date(2026, 9, 15)),
    ("D08", "Project Harbor", "Software-enabled services", "I09", "Broad auction", "Passed", 11.0, 2, date(2026, 8, 15), "Passed - valuation gap", None),
    ("D09", "Project Iris", "Business services", "I10", "Proprietary", "Sourced", 5.5, 4, date(2026, 9, 20), "Request financial summary", date(2026, 10, 12)),
    ("D10", "Project Juniper", "Industrial services", "I07", "Limited auction", "IOI Submitted", 12.5, 4, date(2026, 9, 26), "Prepare management meeting questions", date(2026, 10, 9)),
    ("D11", "Project Kestrel", "Home services", "I04", "Broad auction", "Lost", 8.0, 3, date(2026, 9, 5), "Lost - outbid", None),
    ("D12", "Project Larch", "Specialty distribution", "I03", "Limited auction", "NDA / Teaser", 6.0, 3, date(2026, 9, 29), "Review teaser vs thesis", date(2026, 10, 14)),
]
for r in range(PL0, PL1 + 1):
    row = deals[r - PL0] if r - PL0 < len(deals) else [None] * 11
    for c, v in zip((1, 2, 3, 4, 6, 7, 8, 9, 10, 11, 12), row):   # column E is the intermediary lookup formula
        put(PL, r, c, v, {8: "$M", 9: "n", 10: "d", 12: "d"}.get(c), "in")
    put(PL, r, 5, f'=IF($D{r}="","",IFERROR(INDEX({D}$B${DR0}:$B${DR1},MATCH($D{r},{D}$A${DR0}:$A${DR1},0)),"UNKNOWN ID"))')
    put(PL, r, 13, f'=IF($A{r}="","",IFERROR(VLOOKUP($G{r},{STAGES},3,FALSE),0))', "%")
    put(PL, r, 14, f'=IF($A{r}="","",IFERROR(VLOOKUP($G{r},{STAGES},4,FALSE),0))', "n")
    put(PL, r, 15, f'=IF($A{r}="","",IF(AND($H{r}>={EMIN},$H{r}<={EMAX}),1,IF(AND($H{r}>={EMIN}*0.5,$H{r}<={EMAX}*1.5),0.5,0)))', "n1")
    put(PL, r, 16, f'=IF($A{r}="","",IFERROR(VLOOKUP($F{r},{PROCS},2,FALSE),0))', "n1")
    put(PL, r, 17, f'=IF($A{r}="","",IFERROR(INDEX({D}$H${DR0}:$H${DR1},MATCH($D{r},{D}$A${DR0}:$A${DR1},0)),1))', "n")
    put(PL, r, 18, f'=IF($A{r}="","",ROUND(100*({WST}*M{r}+{WFIT}*($I{r}-1)/4+{WSZ}*O{r}+{WPR}*P{r}+{WREL}*(Q{r}-1)/4),1))', "n1")
    put(PL, r, 19, f'=IF(OR($A{r}="",N{r}<>1),"",R{r}+(1000-ROW())/10^7)', "x")
    put(PL, r, 20, f'=IF(S{r}="","",COUNTIF($S${PL0}:$S${PL1},">"&S{r})+1)', "n")
    put(PL, r, 21, f'=IF(OR($A{r}="",$J{r}=""),"",{ASOF}-$J{r})', "n")
    put(PL, r, 22, f'=IF(OR($A{r}="",N{r}<>1,$L{r}=""),"",MAX(0,{ASOF}-$L{r}))', "n")
    put(PL, r, 23, f'=IF($A{r}="","",IF(N{r}<>1,"Inactive - "&$G{r},IF($L{r}="","NO NEXT ACTION",IF(V{r}>0,"OVERDUE "&V{r}&"d",'
                   f'IF($L{r}-{ASOF}<={WIN},"DUE IN "&($L{r}-{ASOF})&"d",IF(N(U{r})>{STALE},"STALE","On track"))))))')
    put(PL, r, 24, f'=IF($A{r}="","",$H{r}*M{r}*N{r})', "$M")
dv(PL, f"={S}$A$19:$A$28", f"G{PL0}:G{PL1}"); dv(PL, f"={S}$A$31:$A$33", f"F{PL0}:F{PL1}")
dv(PL, f"={D}$A${DR0}:$A${DR1}", f"D{PL0}:D{PL1}")
RED, AMB, GRY = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFE699", "D9D9D9"))
rng = f"W{PL0}:W{PL1}"
PL.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($W{PL0},7)="OVERDUE"'], fill=RED))
PL.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($W{PL0},6)="DUE IN"'], fill=AMB))
PL.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($W{PL0},8)="Inactive"'], fill=GRY))
widths(PL, [8, 16, 22, 11, 13, 15, 20, 10, 8, 12, 32, 12, 10, 8, 9, 9, 11, 11, 12, 10, 11, 9, 20, 12]); PL.freeze_panes = "C4"

# ---------------- Follow-Up Cadence Tracker ----------------
title(FT, "Follow-Up Cadence Tracker", "Section A: intermediary touch cadence by tier. Section B: every deal's next action. All cells are live formulas.", 11)
section(FT, 3, "A. Intermediary touch cadence (one row per Intermediary Directory row)")
header(FT, 4, ["ID", "Code name", "Tier", "Cadence (days)", "Last contact", "Next touch due", "Days until due",
               "Days overdue", "Cadence flag", "Active deals", "Suggested action"])
for i in range(DR1 - DR0 + 1):
    r, dr = 5 + i, DR0 + i
    put(FT, r, 1, f'=IF({D}A{dr}="","",{D}A{dr})')
    put(FT, r, 2, f'=IF(A{r}="","",{D}B{dr})'); put(FT, r, 3, f'=IF(A{r}="","",{D}G{dr})')
    put(FT, r, 4, f'=IF(A{r}="","",IFERROR(VLOOKUP(C{r},{TIERS},2,FALSE),30))', "n")
    put(FT, r, 5, f'=IF(A{r}="","",{D}I{dr})', "d"); put(FT, r, 6, f'=IF(A{r}="","",E{r}+D{r})', "d")
    put(FT, r, 7, f'=IF(A{r}="","",F{r}-{ASOF})', "n"); put(FT, r, 8, f'=IF(A{r}="","",MAX(0,-G{r}))', "n")
    put(FT, r, 9, f'=IF(A{r}="","",IF(H{r}>0,"OVERDUE "&H{r}&"d",IF(G{r}<={WIN},"DUE IN "&G{r}&"d","On track")))')
    put(FT, r, 10, f'=IF(A{r}="","",{D}L{dr})', "n")
    put(FT, r, 11, f'=IF(A{r}="","",IF(H{r}>0,"Call this week",IF(G{r}<={WIN},"Touch by "&TEXT(F{r},"yyyy-mm-dd"),"Next touch "&TEXT(F{r},"yyyy-mm-dd"))))')
FB0 = 33
section(FT, FB0 - 2, "B. Deal next-action follow-ups (one row per Deal Pipeline Log row)")
header(FT, FB0 - 1, ["Deal ID", "Code name", "Intermediary", "Stage", "Next action", "Next action due", "Days overdue", "Follow-up flag"])
for i in range(PL1 - PL0 + 1):
    r, pr = FB0 + i, PL0 + i
    put(FT, r, 1, f'=IF({P}A{pr}="","",{P}A{pr})')
    for c, col in ((2, "B"), (3, "E"), (4, "G"), (5, "K")):
        put(FT, r, c, f'=IF(A{r}="","",{P}{col}{pr})')
    put(FT, r, 6, f'=IF(A{r}="","",IF({P}L{pr}="","",{P}L{pr}))', "d")
    put(FT, r, 7, f'=IF(A{r}="","",{P}V{pr})', "n"); put(FT, r, 8, f'=IF(A{r}="","",{P}W{pr})')
for rg, col in ((f"I5:I{5 + DR1 - DR0}", "I5"), (f"H{FB0}:H{FB0 + PL1 - PL0}", f"H{FB0}")):
    FT.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col[0]}{col[1:]},7)="OVERDUE"'], fill=RED))
    FT.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col[0]}{col[1:]},6)="DUE IN"'], fill=AMB))
widths(FT, [9, 16, 14, 14, 30, 13, 10, 20, 18, 10, 24]); FT.freeze_panes = "A5"

# ---------------- Ranking Dashboard ----------------
TI = f"{T}$I$5:$I${5 + DR1 - DR0}"; TH = f"{T}$H$5:$H${5 + DR1 - DR0}"; TB = f"{T}$B$5:$B${5 + DR1 - DR0}"
def pc(col):
    return f"{P}${col}${PL0}:${col}${PL1}"
title(DB, "Ranking Dashboard", "Live summary: pipeline KPIs, top-ranked active deals, stage funnel, most overdue follow-ups and one dashboard output line.", 6)
header(DB, 3, ["Metric", "Value", "Source"])
kpis = [
    ("As-of date", f"={ASOF}", "d", "Settings B4"),
    ("Intermediaries tracked", f"=COUNTA({D}$A${DR0}:$A${DR1})", "n", "Intermediary Directory A4:A28"),
    ("Deals logged", f"=COUNTA({pc('A')})", "n", "Deal Pipeline Log A4:A43"),
    ("Active deals", f"=SUM({pc('N')})", "n", "Pipeline N"),
    ("Inactive deals (passed / lost / closed)", "=B6-B7", "n", "Logged - active"),
    ("Active deals inside EBITDA band", f"=COUNTIFS({pc('N')},1,{pc('O')},1)", "n", "Pipeline N and O"),
    ("Probability-weighted active EBITDA ($M)", f"=SUM({pc('X')})", "$M", "Pipeline X"),
    ("Overdue deal next actions", f'=COUNTIF({pc("W")},"OVERDUE*")', "n", "Pipeline W"),
    ("Overdue intermediary touches", f'=COUNTIF({TI},"OVERDUE*")', "n", "Tracker I5:I29"),
    ("Follow-ups due within window", f'=COUNTIF({pc("W")},"DUE IN*")+COUNTIF({TI},"DUE IN*")', "n", "Pipeline W + Tracker I"),
    ("Stale active deals / missing next action", f'=COUNTIF({pc("W")},"STALE")+COUNTIF({pc("W")},"NO NEXT ACTION")', "n", "Pipeline W"),
    ("Pipeline health", '=IF(B11+B12=0,"HEALTHY - no overdue follow-ups",IF(B11+B12<=3,"ATTENTION - clear overdue follow-ups this week",'
                        '"AT RISK - "&(B11+B12)&" overdue follow-ups; clear backlog before new outreach"))', None, "B11 + B12 gates"),
]
for i, (lab, f, fm, src) in enumerate(kpis):
    r = 4 + i
    put(DB, r, 1, lab, bold=True); put(DB, r, 2, f, fm, "k" if r == 15 else None, bold=r == 15); put(DB, r, 3, src)
section(DB, 17, "Top 5 active deals by priority score")
header(DB, 18, ["Rank", "Deal", "Priority score", "Stage", "Intermediary", "Follow-up flag"])
for k in range(1, 6):
    r = 18 + k; m = f"MATCH($A{r},{pc('T')},0)"
    put(DB, r, 1, k, "n", bold=True)
    for c, col, fm in ((2, "B", None), (3, "R", "n1"), (4, "G", None), (5, "E", None), (6, "W", None)):
        put(DB, r, c, f'=IFERROR(INDEX({pc(col)},{m}),"-")', fm)
section(DB, 25, "Stage funnel")
header(DB, 26, ["Stage", "Deals (all)", "Active", "EBITDA ($M, all)", "Prob.-weighted EBITDA ($M)"])
for i in range(10):
    r = 27 + i
    put(DB, r, 1, f"={S}A{19 + i}", bold=True)
    put(DB, r, 2, f"=COUNTIF({pc('G')},$A{r})", "n"); put(DB, r, 3, f"=COUNTIFS({pc('G')},$A{r},{pc('N')},1)", "n")
    put(DB, r, 4, f"=SUMIF({pc('G')},$A{r},{pc('H')})", "$M"); put(DB, r, 5, f"=SUMIF({pc('G')},$A{r},{pc('X')})", "$M")
put(DB, 37, 1, "Total", bold=True)
for c in "BCDE":
    put(DB, 37, ord(c) - 64, f"=SUM({c}27:{c}36)", "n" if c in "BC" else "$M", bold=True)
section(DB, 39, "Most overdue follow-ups")
header(DB, 40, ["Item", "Name", "Days overdue", "Action"])
put(DB, 41, 1, "Most overdue deal action", bold=True)
put(DB, 41, 2, f'=IF(MAX({pc("V")})=0,"None",INDEX({pc("B")},MATCH(MAX({pc("V")}),{pc("V")},0)))')
put(DB, 41, 3, f'=MAX({pc("V")})', "n")
put(DB, 41, 4, f'=IF(C41=0,"-",INDEX({pc("K")},MATCH(C41,{pc("V")},0)))')
put(DB, 42, 1, "Most overdue intermediary touch", bold=True)
put(DB, 42, 2, f'=IF(MAX({TH})=0,"None",INDEX({TB},MATCH(MAX({TH}),{TH},0)))')
put(DB, 42, 3, f"=MAX({TH})", "n"); put(DB, 42, 4, '=IF(C42=0,"-","Call this week; confirm active mandates")')
put(DB, 44, 1, "DASHBOARD OUTPUT", bold=True)
put(DB, 44, 2, '="Top priority: "&B19&" (score "&IFERROR(TEXT(C19,"0.0"),"-")&", "&D19&", via "&E19&"). "&B7&" active of "&B6&" deals; '
               'probability-weighted EBITDA $"&TEXT(B10,"0.00")&"M. Follow-ups: "&B11&" overdue deal actions, "&B12&" overdue intermediary touches, "'
               f'&B13&" due within "&{WIN}&" days. Most overdue: "&B41&" ("&C41&"d) and "&B42&" ("&C42&"d). Health: "&B15&"."', kind="k", bold=True)
put(DB, 45, 1, "Pipeline health formula (FORMULATEXT of B15)", bold=True); put(DB, 45, 2, "=_xlfn.FORMULATEXT(B15)")
DB.merge_cells("B44:F44"); DB.merge_cells("B45:F45"); DB.row_dimensions[44].height = 78; DB.row_dimensions[45].height = 48
DB.row_dimensions[15].height = 46
widths(DB, [40, 34, 16, 20, 22, 20]); DB.freeze_panes = "A4"

# ---------------- Anonymized Worked Example ----------------
title(WE, "Anonymized Worked Example - Sponsor X lower-middle-market pipeline", "Code names, regions and scaled figures only. The dataset lives in the Intermediary Directory (rows 4-13) and Deal Pipeline Log (rows 4-15); this tab reads the live outputs.", 6)
section(WE, 3, "Scenario")
WE["A4"] = ("Sponsor X targets $3-15M EBITDA services and distribution businesses. It tracks 10 intermediaries and 12 opportunities. "
            "The question for the Monday pipeline meeting is: which deals deserve time this week, and which relationships are slipping?")
section(WE, 6, "Live snapshot (links to Ranking Dashboard)")
header(WE, 7, ["Metric", "Value"])
snap = [("As-of date", f"={B}B4", "d"), ("Deals logged / active", f'={B}B6&" / "&{B}B7', None),
        ("Probability-weighted active EBITDA ($M)", f"={B}B10", "$M"), ("Overdue deal next actions", f"={B}B11", "n"),
        ("Overdue intermediary touches", f"={B}B12", "n"), ("Pipeline health", f"={B}B15", None),
        ("Rank 1", f'={B}B19&" - score "&TEXT({B}C19,"0.0")', None), ("Rank 2", f'={B}B20&" - score "&TEXT({B}C20,"0.0")', None),
        ("Rank 3", f'={B}B21&" - score "&TEXT({B}C21,"0.0")', None)]
for i, (lab, f, fm) in enumerate(snap):
    put(WE, 8 + i, 1, lab, bold=True); put(WE, 8 + i, 2, f, fm)
section(WE, 18, "Named judgment calls")
header(WE, 19, ["ID", "Judgment call", "Choice made", "Alternative rejected", "Live impact", "Impact unit", "Rationale"])
jcs = [
    ("JC1", "Project Ember is above the EBITDA band", "Keep it with half size credit (0.5)", "Exclude it (size fit 0)",
     f"=100*{WSZ}*0.5", "priority points kept", "At $18.0M it is within 50% of the $15M ceiling and comes from an A-tier advisor; keep it visible but ranked below in-band deals."),
    ("JC2", "Project Birch thesis fit", "Fit 3 of 5", "Fit 4 of 5 (size and sector match)",
     f"=100*{WFIT}*0.25", "priority points withheld", "Broad auction with an ask above our underwriting view; sector match alone does not earn a 4."),
    ("JC3", "Passed and lost deals stay in the log", "Keep for intermediary attribution, exclude from ranking", "Delete them",
     f"=COUNTIF({pc('N')},0)", "deals excluded from ranking", "Deals-sourced counts show which intermediaries send relevant flow, even when we pass."),
    ("JC4", "Fixed as-of date instead of TODAY()", "As-of 2026-10-04", "Live TODAY() date",
     f"={B}B11+{B}B12", "overdue items frozen", "A fixed date keeps the example reproducible for validators; live users switch Settings B4 to =TODAY()."),
]
for i, row in enumerate(jcs):
    for c, v in enumerate(row, 1):
        put(WE, 20 + i, c, v, ("n1" if i < 2 else "n") if c == 5 else None, bold=c == 1)
section(WE, 25, "Dashboard output")
put(WE, 26, 1, f"={B}B44", kind="k", bold=True); WE.merge_cells("A26:G26"); WE.row_dimensions[26].height = 64
section(WE, 28, "How to read this example")
for i, t in enumerate(["1. Project Falcon ranks first: proprietary, in-band, at LOI, sourced by a strength-5 operator relationship.",
                       "2. The two most overdue deal actions are early-stage proprietary deals (Granite 19d, Cedar 8d); the most lapsed relationships are Banker I (66d) and CPA F (21d). These are this week's calls.",
                       "3. Change any yellow input (stage, due date, weight) and the rank, flags, dashboard and this tab all recalculate."]):
    WE.cell(29 + i, 1, t)
WE.row_dimensions[13].height = 32
widths(WE, [38, 44, 30, 30, 11, 22, 70])

# ---------------- ReadMe - Validator Instructions ----------------
title(RM, "Deal Flow Pipeline Tracker Toolkit - v1.0", "Post Fiat Analytical Infrastructure Library | LMM private equity deal sourcing | Formula-driven, anonymized", 5)
section(RM, 4, "Purpose")
RM["A5"] = ("Tracks bankers, brokers and other intermediaries, ranks a live deal pipeline with a transparent 0-100 priority score, "
            "and flags overdue deal next actions and lapsed intermediary touch cadences on one dashboard.")
section(RM, 7, "How to use")
for i, t in enumerate(["1. Set the as-of date, EBITDA band, weights, stage probabilities and cadences on Settings and Scoring Rules (yellow cells).",
                       "2. Add intermediaries on Intermediary Directory A4:J28 and deals on Deal Pipeline Log A4:L43 (yellow cells; dropdowns validate stage, process, tier and ID).",
                       "3. Work the Follow-Up Cadence Tracker: every OVERDUE row is a call or email this week.",
                       "4. Read Ranking Dashboard B44 for the one-line pipeline output and B19:F23 for the top 5 deals."]):
    RM.cell(8 + i, 1, t)
section(RM, 13, "Tabs")
header(RM, 14, ["#", "Tab", "What it does"])
desc = ["Purpose, usage, validator checklist and formula reference", "Weights, stage probabilities, process scores, tier cadences, thresholds",
        "Anonymized intermediary records plus live deal roll-ups", "Deal log with stage, next action, priority score, rank and follow-up flag",
        "Intermediary cadence and deal next-action follow-ups with overdue flags", "KPIs, top 5 ranked deals, stage funnel, most overdue, dashboard output",
        "Sponsor X example: live snapshot, 4 named judgment calls, dashboard output"]
for i, (n, t) in enumerate(zip(N, desc)):
    put(RM, 15 + i, 1, i + 1); put(RM, 15 + i, 2, n, bold=True); put(RM, 15 + i, 3, t)
RM["A23"] = "Colour legend: yellow = input, blue = formula, green = key output. Tab 1 is 'ReadMe - Validator Instructions' because Excel does not allow '/' in tab names."
section(RM, 25, "Validator checklist (criterion -> tab -> cell range -> live check)")
header(RM, 26, ["#", "Criterion", "Tab", "Cell range", "Live check"])
chk = [
    ("Validator Read Me", N[0], "A1:E45", '="Tabs in workbook: "&_xlfn.SHEETS()'),
    ("Anonymized intermediary directory", N[2], "A4:M28", f'=COUNTA({D}A4:A28)&" intermediaries, code names only"'),
    ("Deal pipeline log with stage and next action", N[3], "A4:X43 (stage G, next action K/L)", f'=COUNTA({pc("A")})&" deals; "&COUNTA({pc("K")})&" with a next action"'),
    ("Follow-up cadence tracker with overdue flags", N[4], "A5:K29 and A33:H72", f'={B}B12&" intermediary + "&{B}B11&" deal follow-ups overdue"'),
    ("Live priority score and rank", N[3], "R4:T43; weights Settings B11:C16", f'="Rank 1: "&{B}B19&" ("&TEXT({B}C19,"0.0")&")"'),
    ("Weights sum to 100%", N[1], "B16:C16", f"={S}C16"),
    ("Summary dashboard", N[5], "A3:F45; output B44", f"={B}B15"),
    ("Anonymized worked example", N[6], "A1:G31 (4 judgment calls A20:G23)", f'=COUNTA({W}A20:A23)&" judgment calls"'),
    ("Visible live formulas", N[0], "Formula reference A39:C43 (FORMULATEXT)", '=COUNTA(C39:C43)&" key formulas shown as live FORMULATEXT"'),
    ("Zero formula errors", "All tabs", "Excel SpecialCells(xlFormulas, xlErrors) scan", "See README.md (recalculated in Excel)"),
]
for i, row in enumerate(chk):
    r = 27 + i
    put(RM, r, 1, i + 1)
    for c, v in enumerate(row, 2):
        put(RM, r, c, v, bold=c == 2)
section(RM, 38, "Formula reference (live FORMULATEXT of key cells)")
for i, (lab, ref) in enumerate([("Priority score (Deal Pipeline Log R4)", f"{P}R4"), ("Priority rank (Deal Pipeline Log T4)", f"{P}T4"),
                                ("Follow-up flag (Deal Pipeline Log W4)", f"{P}W4"), ("Cadence flag (Follow-Up Cadence Tracker I5)", f"{T}I5"),
                                ("Pipeline health (Ranking Dashboard B15)", f"{B}B15")]):
    r = 39 + i
    put(RM, r, 1, i + 1); put(RM, r, 2, lab, bold=True); put(RM, r, 3, f"=_xlfn.FORMULATEXT({ref})"); RM.merge_cells(f"C{r}:E{r}")
    RM.row_dimensions[r].height = 60
widths(RM, [6, 44, 34, 40, 46])

wb.save(OUT)
print("saved", OUT)




